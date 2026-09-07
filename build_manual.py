from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
OUT=r'C:\\xampp\\htdocs\\Taxguard\\TaxGuard_User_Manual.docx'
d=Document(); s=d.sections[0]; s.top_margin=Inches(.7); s.bottom_margin=Inches(.65); s.left_margin=Inches(.8); s.right_margin=Inches(.8)
for n,sz in [('Title',31),('Heading 1',19),('Heading 2',13)]:
 st=d.styles[n]; st.font.name='Aptos'; st.font.size=Pt(sz); st.font.bold=True; st.font.color.rgb=RGBColor(11,39,67)
st=d.styles['Normal']; st.font.name='Aptos'; st.font.size=Pt(10); st.font.color.rgb=RGBColor(35,48,65); st.paragraph_format.space_after=Pt(6); st.paragraph_format.line_spacing=1.08
def shade(c,f):
 x=c._tc.get_or_add_tcPr(); z=OxmlElement('w:shd'); z.set(qn('w:fill'),f); x.append(z)
def table(h,rows):
 t=d.add_table(rows=1,cols=len(h)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for i,x in enumerate(h):
  c=t.rows[0].cells[i]; c.text=x; shade(c,'123B5D')
  for r in c.paragraphs[0].runs:r.font.bold=True;r.font.color.rgb=RGBColor(255,255,255);r.font.size=Pt(9)
 for j,row in enumerate(rows):
  cs=t.add_row().cells
  for i,x in enumerate(row):
   cs[i].text=str(x)
   if j%2:shade(cs[i],'F3F7FB')
   for p in cs[i].paragraphs:
    for r in p.runs:r.font.size=Pt(9)
 d.add_paragraph()
def num(x): d.add_paragraph(x,style='List Number')
def bullet(x): d.add_paragraph(x,style='List Bullet')
p=d.add_paragraph(style='Title');p.add_run('TaxGuard User Manual')
p=d.add_paragraph();r=p.add_run('Compliance workspace operations and local SQLite storage');r.font.size=Pt(15);r.font.color.rgb=RGBColor(39,102,219)
d.add_paragraph('Version 1.0  |  September 2026')
p=d.add_paragraph();p.add_run('Purpose. ').bold=True;p.add_run('This manual explains the current TaxGuard demo, its client and filing workflow, and its shared local database.')
d.add_paragraph('TaxGuard organizes obligations and records submissions made through external filing channels. It does not submit tax returns to BIR.')
table(['At a glance','Current behavior'],[('Audience','One company team; no connected user-account screen yet.'),('Storage','SQLite shared by localhost and Electron.'),('Access','http://localhost/Taxguard/ or the Electron app.'),('Data included','Six sample clients and 133 sample filings for 2024–2026.'),('Deadline source','Supplied schedule workbook data; verify before production use.')])
d.add_heading('How the system fits together',1);d.add_paragraph('A client is assigned forms. Each form has periods and due dates. TaxGuard expands those assignments into yearly obligations. When a submission is recorded, the matching obligation becomes complete and appears in dashboard summaries and recent filings.')
d.add_page_break();d.add_heading('1  Start the system',1);d.add_heading('Local browser version',2)
for x in ['Start Apache in XAMPP.','Open http://localhost/Taxguard/.','Check the footer: “Saved to SQLite on this computer” confirms the connection.']:num(x)
d.add_paragraph('The browser sends local JSON requests to api.php. PHP runs the same database validation and transaction code used by Electron, so both interfaces read the same records.')
d.add_heading('Electron version',2)
for x in ['Open PowerShell in the Taxguard folder.','Run npm.cmd install once, then npm.cmd start.','Use the same navigation and forms as the browser.']:num(x)
d.add_paragraph('Returning focus to a window refreshes SQLite records when no modal is open. A stale save is rejected instead of overwriting a change made elsewhere.')
d.add_heading('2  Navigation',1)
table(['Area','Use it to'],[('Overview','See counts, progress, completion, and recent filings.'),('Client directory','Add, search, review, and edit master client records.'),('Compliance tracker','Review yearly obligations and record filings.'),('Deadline reference','Review, add, and edit form schedules.'),('Settings','Export records and import a browser export into an empty desktop database.')])
d.add_page_break();d.add_heading('3  Add and maintain clients',1);d.add_paragraph('The client record is the source for tracker obligations. Complete the required fields before saving.')
table(['Field','What to enter'],[('Client name','Legal or working name.'),('TIN','Exactly 000-000-000-000; it must be unique.'),('Start of filing','Date from which obligations begin.'),('Business type','Sole proprietorship, Partnership, Corporation, and so on.'),('Tax type','VAT or NVAT.'),('Client status','Active, Inactive, Closed, or For closure.'),('Required forms','Forms that apply to this client.'),('Remarks','Optional internal notes.')])
d.add_paragraph('Invalid dates, duplicate TINs, unknown forms, and incomplete fields are rejected before commit; the previous database state remains unchanged.')
d.add_heading('Required forms',2);d.add_paragraph('Choose forms through the picker. Selected forms appear as pills. The assignment controls which periods are generated in the tracker.')
d.add_heading('4  Manage deadline references',1);d.add_paragraph('Open a form to view its periods and due dates. Use Add deadline to create a reference, or Edit schedule to change an existing one. Existing values are loaded into the editor so they are not cleared accidentally. Verify dates against current filing guidance before production use.')
d.add_page_break();d.add_heading('5  Record a filing',1);d.add_paragraph('After filing through the normal external channel, record the result in Compliance tracker.')
for x in ['Select the tax year and locate the client, form, and period.','Choose Record filing, or View filing for an existing submission.','Enter the filing date.','Enter the confirmation or reference number.','Add optional remarks and choose Save filing.']:num(x)
d.add_paragraph('The current interface treats the confirmation/reference number as required. Without it, the filing will not save.')
d.add_heading('Status meanings',2)
table(['Status','Meaning'],[('Incomplete','No filing recorded.'),('Overdue','Incomplete and past the demo comparison date.'),('Complete','A filing date and saved record exist.'),('Client annual status','Summary of all applicable obligations for that client and year.')])
d.add_heading('6  Dashboard',1);d.add_paragraph('Overview summarizes the selected year: completed filings, remaining obligations, overdue items, progress by form, portfolio completion, and recent filings. Clicking a recent filing opens Client directory with that client filtered and its detail modal open.')
d.add_page_break();d.add_heading('7  Data and backups',1);d.add_heading('Shared database',2);d.add_paragraph('The active SQLite file is database/taxguard.db. Both localhost and Electron use it. Transactions ensure a failed save does not leave a half-written change.')
d.add_heading('Export and import',2)
for x in ['Open Settings and choose Export records.','Keep the downloaded JSON private because it contains client records.','In the desktop app, choose Import browser records while the desktop client list is empty.','Select the export; validation completes before the import is committed.']:num(x)
d.add_paragraph('GitHub Pages cannot access this local database and remains a separate browser-storage demo.')
d.add_heading('Current limitations',2)
for x in ['Six sample clients and 133 sample filings are seeded once for 2024–2026.','Sample deadlines come from supplied workbook schedules and are not verified current BIR deadlines.','Overdue calculations use the demo date September 5, 2026.','No connected company login or user-account workflow exists yet.','The filing reference number is currently required.']:bullet(x)
d.add_heading('8  Troubleshooting',1)
table(['Symptom','Check'],[('No browser records','Start Apache, confirm localhost URL, refresh.'),('SQLite connection failed','Check Node.js at C:\\Program Files\\nodejs\\node.exe or configure TAXGUARD_NODE_PATH.'),('Electron shows old data','Close and restart Electron.'),('Records changed in another window','Refresh, reopen the record, and save again.'),('Import unavailable','The desktop database already has clients; use the safety check as designed.')])
d.add_paragraph('TaxGuard · Integrated multi-year compliance monitoring').alignment=WD_ALIGN_PARAGRAPH.CENTER
d.save(OUT);print(OUT)
